from pathlib import Path

from leads.export import to_csv, to_xlsx
from leads.models import Ad, Lead


def sample_leads():
    lead_a = Lead(page_id="p1", page_name="Top Lead", score=80, tier="A")
    lead_a.ads["1"] = Ad(ad_id="1", page_id="p1", page_name="Top Lead", keyword="k",
                          sector="s", source="api")
    lead_b = Lead(page_id="p2", page_name="Excluded", excluded="blocklist", score=0, tier="excluido")
    return {"p1": lead_a, "p2": lead_b}


def test_to_csv_excludes_blocked_and_writes_header(tmp_path):
    path = tmp_path / "out.csv"
    n = to_csv(sample_leads(), path)
    assert n == 1
    content = path.read_text(encoding="utf-8")
    assert "Top Lead" in content
    assert "Excluded" not in content


def test_to_xlsx_creates_file_with_excluded_sheet(tmp_path):
    path = tmp_path / "out.xlsx"
    n = to_xlsx(sample_leads(), path)
    assert n == 1
    assert path.exists()

    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert "Leads" in wb.sheetnames
    assert "Excluidos" in wb.sheetnames
