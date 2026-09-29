from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import Lead

COLUMNS = [
    ("Tier", lambda l: l.tier),
    ("Puntuación", lambda l: l.score),
    ("Página", lambda l: l.page_name),
    ("Sectores", lambda l: ", ".join(sorted(l.sectors))),
    ("Web", lambda l: l.website),
    ("Instagram", lambda l: f"instagram.com/{l.instagram}" if l.instagram else ""),
    ("TikTok", lambda l: f"tiktok.com/@{l.tiktok}" if l.tiktok else ""),
    ("Email", lambda l: ", ".join(l.emails)),
    ("Teléfono", lambda l: ", ".join(l.phones)),
    ("WhatsApp", lambda l: f"https://wa.me/{l.whatsapp}" if l.whatsapp else ""),
    ("Anuncios activos", lambda l: l.active_ads),
    ("Creatividades distintas", lambda l: l.distinct_creatives),
    ("Días de campaña", lambda l: l.days_running()),
    ("Alcance UE (est.)", lambda l: l.eu_reach or ""),
    ("Plataformas", lambda l: ", ".join(l.platforms)),
    ("Stack web", lambda l: l.platform),
    ("Pixel Meta", lambda l: "Sí" if l.has_meta_pixel else ""),
    ("Ejemplo de anuncio", lambda l: l.sample_ad),
    ("Biblioteca de Anuncios", lambda l: l.ad_library_url),
    ("Motivos de la puntuación", lambda l: " · ".join(l.reasons)),
]

TIER_FILL = {
    "A": PatternFill("solid", fgColor="C6EFCE"),
    "B": PatternFill("solid", fgColor="FFEB9C"),
    "C": PatternFill("solid", fgColor="F2F2F2"),
}


def _sorted(leads: dict[str, Lead]) -> list[Lead]:
    included = [l for l in leads.values() if not l.excluded]
    return sorted(included, key=lambda l: l.score, reverse=True)


def to_csv(leads: dict[str, Lead], path: Path) -> int:
    rows = _sorted(leads)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([c[0] for c in COLUMNS])
        for lead in rows:
            w.writerow([fn(lead) for _, fn in COLUMNS])
    return len(rows)


def to_xlsx(leads: dict[str, Lead], path: Path) -> int:
    rows = _sorted(leads)
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([c[0] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F1F1F")
        cell.alignment = Alignment(horizontal="center")

    for lead in rows:
        ws.append([fn(lead) for _, fn in COLUMNS])
        fill = TIER_FILL.get(lead.tier)
        if fill:
            ws.cell(row=ws.max_row, column=1).fill = fill

    for i, (name, _) in enumerate(COLUMNS, start=1):
        width = 40 if name in ("Ejemplo de anuncio", "Motivos de la puntuación", "Biblioteca de Anuncios") else 18
        ws.column_dimensions[get_column_letter(i)].width = width

    if rows:
        ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"
        table = Table(displayName="Leads", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "A2"

    excluded = [l for l in leads.values() if l.excluded]
    if excluded:
        ws2 = wb.create_sheet("Excluidos")
        ws2.append(["Página", "Motivo de exclusión", "Web"])
        for l in excluded:
            ws2.append([l.page_name, l.excluded, l.website])

    wb.save(path)
    return len(rows)
